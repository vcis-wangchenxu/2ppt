#!/usr/bin/env python3
"""Summarize a PPTX deck for humans or downstream automation."""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from pathlib import Path
from typing import Any, Iterable
from xml.etree import ElementTree as ET

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError as exc:  # pragma: no cover
    raise SystemExit(
        "缺少 python-pptx 运行依赖；请运行：pip install -r requirements.txt"
    ) from exc

from runtime_common import UserInputError
from validate_pptx import validate as validate_package


EMU_PER_INCH = 914400
NS_REL = "http://schemas.openxmlformats.org/package/2006/relationships"


def _inches(value: int) -> float:
    return round(value / EMU_PER_INCH, 3)


def _shape_kind(shape: Any) -> str:
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "has_table", False):
        return "table"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "image"
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if shape.shape_type == MSO_SHAPE_TYPE.LINE:
        return "line"
    if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
        return "text"
    if shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
        return "shape"
    return str(shape.shape_type).split(" ", 1)[0].lower()


def _alt_text(shape: Any) -> str:
    element = shape._element
    for path in ("nvPicPr", "nvSpPr", "nvCxnSpPr", "nvGraphicFramePr", "nvGrpSpPr"):
        non_visual = getattr(element, path, None)
        props = getattr(non_visual, "cNvPr", None) if non_visual is not None else None
        if props is not None:
            return props.get("descr", "")
    return ""


def _shape_summary(shape: Any) -> dict[str, Any]:
    item: dict[str, Any] = {
        "id": shape.shape_id,
        "name": shape.name,
        "type": _shape_kind(shape),
        "bounds": {
            "x": _inches(shape.left),
            "y": _inches(shape.top),
            "w": _inches(shape.width),
            "h": _inches(shape.height),
        },
    }
    alt = _alt_text(shape)
    if alt:
        item["alt"] = alt
    if getattr(shape, "has_text_frame", False):
        text = shape.text.strip()
        if text:
            item["text"] = text[:1000]
    if getattr(shape, "has_table", False):
        item["rows"] = len(shape.table.rows)
        item["columns"] = len(shape.table.columns)
    if getattr(shape, "has_chart", False):
        chart = shape.chart
        item["chart_type"] = str(chart.chart_type)
        item["series"] = [str(series.name) for series in chart.series]
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        item["children"] = [_shape_summary(child) for child in shape.shapes]
    return item


def _flatten(items: Iterable[dict[str, Any]]) -> Iterable[dict[str, Any]]:
    for item in items:
        yield item
        yield from _flatten(item.get("children", []))


def _package_summary(path: Path) -> dict[str, int]:
    with zipfile.ZipFile(path, "r") as archive:
        media = charts = embeddings = external = 0
        for item in archive.infolist():
            name = item.filename
            if name.startswith("ppt/media/") and not item.is_dir():
                media += 1
            elif name.startswith("ppt/charts/chart") and name.endswith(".xml"):
                charts += 1
            elif name.startswith("ppt/embeddings/") and not item.is_dir():
                embeddings += 1
            if name.endswith(".rels"):
                try:
                    root = ET.fromstring(archive.read(item))
                except ET.ParseError:
                    continue
                external += sum(
                    1
                    for relationship in root.findall(f"{{{NS_REL}}}Relationship")
                    if relationship.get("TargetMode") == "External"
                )
        return {
            "members": len([item for item in archive.infolist() if not item.is_dir()]),
            "media": media,
            "charts": charts,
            "embeddings": embeddings,
            "external_relationships": external,
        }


def inspect(path: Path) -> dict[str, Any]:
    resolved = path.expanduser().resolve()
    if not resolved.is_file():
        raise UserInputError(f"文件不存在：{resolved}")
    preflight = validate_package(resolved)
    if not preflight["ok"]:
        summary = "；".join(preflight["errors"][:3])
        raise UserInputError(f"PPTX 安全/结构预检未通过：{summary}")
    try:
        presentation = Presentation(resolved)
    except Exception as exc:
        raise UserInputError(f"python-pptx 无法打开文件：{exc}") from exc
    properties = presentation.core_properties
    report: dict[str, Any] = {
        "ok": True,
        "path": str(resolved),
        "bytes": resolved.stat().st_size,
        "slide_size": {
            "width": _inches(presentation.slide_width),
            "height": _inches(presentation.slide_height),
        },
        "slide_count": len(presentation.slides),
        "properties": {
            key: getattr(properties, key) or ""
            for key in ("title", "subject", "author", "keywords", "comments", "language")
        },
        "package": _package_summary(resolved),
        "validation_warnings": preflight["warnings"],
        "slides": [],
    }
    totals = {kind: 0 for kind in ("text", "shape", "line", "image", "table", "chart", "group", "other")}
    for slide_index, slide in enumerate(presentation.slides, start=1):
        try:
            shapes = [_shape_summary(shape) for shape in slide.shapes]
            flat = list(_flatten(shapes))
            counts: dict[str, int] = {}
            for item in flat:
                kind = item["type"] if item["type"] in totals else "other"
                counts[kind] = counts.get(kind, 0) + 1
                totals[kind] += 1
            text_candidates = [item.get("text", "").strip() for item in flat]
            title = next((text.splitlines()[0][:200] for text in text_candidates if text), "")
            notes = ""
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            slide_name = slide._element.cSld.get("name", "")
        except Exception as exc:
            raise UserInputError(
                f"无法解析第 {slide_index} 页的对象：{type(exc).__name__}: {exc}"
            ) from exc
        report["slides"].append(
            {
                "number": slide_index,
                "name": slide_name,
                "title": title,
                "notes": notes[:5000],
                "shape_count": len(flat),
                "counts": counts,
                "shapes": shapes,
            }
        )
    report["elements"] = totals
    return report


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="检查 PPTX 的页数、元素、文本和包内容")
    parser.add_argument("file", type=Path, help="要检查的 .pptx 文件")
    parser.add_argument("--json", action="store_true", help="输出完整机器可读 JSON")
    args = parser.parse_args(argv)
    try:
        report = inspect(args.file)
    except (UserInputError, OSError, zipfile.BadZipFile) as exc:
        parser.exit(1, f"检查失败：{exc}\n")
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True, default=str))
    else:
        print(f"PPTX：{report['path']}")
        print(
            f"幻灯片：{report['slide_count']}；画布："
            f"{report['slide_size']['width']}×{report['slide_size']['height']} 英寸"
        )
        print(
            "元素："
            + "，".join(f"{kind}={count}" for kind, count in report["elements"].items() if count)
        )
        for slide in report["slides"]:
            title = slide["title"] or "（无标题）"
            print(f"{slide['number']:>3}. {title}；元素 {slide['shape_count']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
