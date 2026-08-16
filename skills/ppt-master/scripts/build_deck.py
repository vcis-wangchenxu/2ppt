#!/usr/bin/env python3
"""Build a native, editable PowerPoint deck from a small JSON specification."""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import sys
import tempfile
import warnings
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Mapping

try:
    from PIL import Image as PILImage
    from pptx import Presentation
    from pptx.chart.data import ChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
except ImportError as exc:  # pragma: no cover - exercised by doctor.py instead
    raise SystemExit(
        "缺少 python-pptx 或 Pillow 运行依赖；请运行：pip install -r requirements.txt"
    ) from exc

from runtime_common import (
    SCHEMA_VERSION,
    UserInputError,
    boolean,
    choice,
    color,
    integer,
    load_json_object,
    number,
    optional_string,
    require_list,
    require_mapping,
    safe_relative_file,
)


DEFAULT_THEME = {
    "fonts": {"heading": "Aptos Display", "body": "Aptos"},
    "colors": {
        "background": "F7F8FA",
        "text": "172033",
        "primary": "2563EB",
        "secondary": "0F766E",
        "accent": "F59E0B",
        "muted": "667085",
        "white": "FFFFFF",
    },
}

SLIDE_SIZES = {"wide": (13.333333, 7.5), "standard": (10.0, 7.5)}
ROUTES = ("generate", "create-template", "fill-template", "enhance")
BUILDER_ROUTES = ("generate", "create-template")
PROFILES = ("ordinary", "beautify", "image-to-pptx")

TEXT_STYLE_KEYS = {"font_face", "font_size", "color", "bold", "italic"}
TEXT_FRAME_KEYS = {"text", "paragraphs", "align", "valign", "margin", *TEXT_STYLE_KEYS}
COMMON_KEYS = {"type", "name", "alt"}
RECT_KEYS = {"x", "y", "w", "h"}
ELEMENT_KEYS = {
    "text": COMMON_KEYS | RECT_KEYS | TEXT_FRAME_KEYS | {"fill", "line"},
    "shape": COMMON_KEYS
    | RECT_KEYS
    | TEXT_FRAME_KEYS
    | {"shape_type", "fill", "line", "rotation"},
    "line": COMMON_KEYS
    | {"x1", "y1", "x2", "y2", "color", "width", "dash", "begin_arrow", "end_arrow"},
    "image": COMMON_KEYS | RECT_KEYS | {"path", "fit"},
    "table": COMMON_KEYS
    | RECT_KEYS
    | {
        "rows",
        "header_rows",
        "column_widths",
        "align",
        "font_face",
        "font_size",
        "color",
        "header_color",
        "header_fill",
        "cell_fill",
    },
    "chart": COMMON_KEYS
    | RECT_KEYS
    | {"chart", "categories", "series", "title", "legend", "colors", "style", "data_labels"},
}
PARAGRAPH_KEYS = {"text", "runs", "align", "level"}
RUN_KEYS = {"text", *TEXT_STYLE_KEYS}
LINE_STYLE_KEYS = {"color", "width", "dash"}
CHART_SERIES_KEYS = {"name", "values"}
MAX_PARAGRAPHS = 200
MAX_RUNS_PER_PARAGRAPH = 200
MAX_CHART_SERIES = 100
MAX_IMAGE_BYTES = 100 * 1024 * 1024
MAX_IMAGE_PIXELS = 80_000_000

SHAPES = {
    "rect": MSO_AUTO_SHAPE_TYPE.RECTANGLE,
    "rectangle": MSO_AUTO_SHAPE_TYPE.RECTANGLE,
    "rounded_rect": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
    "rounded_rectangle": MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
    "ellipse": MSO_AUTO_SHAPE_TYPE.OVAL,
    "oval": MSO_AUTO_SHAPE_TYPE.OVAL,
    "triangle": MSO_AUTO_SHAPE_TYPE.ISOSCELES_TRIANGLE,
    "diamond": MSO_AUTO_SHAPE_TYPE.DIAMOND,
    "chevron": MSO_AUTO_SHAPE_TYPE.CHEVRON,
    "pentagon": MSO_AUTO_SHAPE_TYPE.PENTAGON,
    "hexagon": MSO_AUTO_SHAPE_TYPE.HEXAGON,
    "arrow_right": MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
    "arrow_left": MSO_AUTO_SHAPE_TYPE.LEFT_ARROW,
    "arrow_up": MSO_AUTO_SHAPE_TYPE.UP_ARROW,
    "arrow_down": MSO_AUTO_SHAPE_TYPE.DOWN_ARROW,
}

ALIGNMENTS = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}

VERTICAL_ALIGNMENTS = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}

DASHES = {
    "solid": MSO_LINE_DASH_STYLE.SOLID,
    "dash": MSO_LINE_DASH_STYLE.DASH,
    "dot": MSO_LINE_DASH_STYLE.ROUND_DOT,
    "dash_dot": MSO_LINE_DASH_STYLE.DASH_DOT,
    "long_dash": MSO_LINE_DASH_STYLE.LONG_DASH,
}

CHARTS = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
}

LEGENDS = {
    "bottom": XL_LEGEND_POSITION.BOTTOM,
    "top": XL_LEGEND_POSITION.TOP,
    "left": XL_LEGEND_POSITION.LEFT,
    "right": XL_LEGEND_POSITION.RIGHT,
}

FIXED_TIMESTAMP = (2000, 1, 1, 0, 0, 0)
CORE_DATE_RE = re.compile(
    rb"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)"
)
CHART_AXIS_ID_RE = re.compile(
    rb'(<c:(?:axId|crossAx)\b[^>]*\bval=")(-\d+)(")'
)


def _normalize_chart_axis_ids(payload: bytes) -> bytes:
    """Convert python-pptx signed chart axis IDs to valid unsigned values.

    python-pptx may serialize randomly generated 32-bit axis identifiers as
    signed decimal strings. PowerPoint tolerates those values, but strict
    Open XML readers correctly treat ``c:axId`` and ``c:crossAx`` as unsigned
    integers and reject the package. Mapping the signed representation modulo
    2**32 preserves every reference while restoring schema compatibility.
    """

    def replace(match: re.Match[bytes]) -> bytes:
        signed = int(match.group(2))
        unsigned = signed % (1 << 32)
        return match.group(1) + str(unsigned).encode("ascii") + match.group(3)

    return CHART_AXIS_ID_RE.sub(replace, payload)


def _unknown_keys(value: Mapping[str, Any], allowed: set[str], field: str) -> None:
    extras = sorted(set(value) - allowed)
    if extras:
        raise UserInputError(f"{field} 包含未知字段：{', '.join(extras)}")


def _font_name(value: Any, field: str) -> str:
    result = optional_string(value, field, maximum=200).strip()
    if not result:
        raise UserInputError(f"{field} 不能为空")
    return result


def _plain_cell(value: Any, field: str) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (str, int, float)):
        text = str(value)
        if len(text) > 5000:
            raise UserInputError(f"{field} 过长")
        optional_string(text, field, maximum=5000)
        return text
    raise UserInputError(f"{field} 只能是字符串、数字、布尔值或 null")


def _set_east_asian_font(run: Any, face: str) -> None:
    """PowerPoint stores East Asian typefaces separately from Latin ones."""
    rpr = run._r.get_or_add_rPr()
    ea = rpr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = OxmlElement("a:ea")
        rpr.append(ea)
    ea.set("typeface", face)


def _set_alt_text(shape: Any, alt: str | None, name: str | None = None) -> None:
    if not alt and not name:
        return
    element = shape._element
    for path in ("nvPicPr", "nvSpPr", "nvCxnSpPr", "nvGraphicFramePr"):
        non_visual = getattr(element, path, None)
        if non_visual is None:
            continue
        props = getattr(non_visual, "cNvPr", None)
        if props is None:
            continue
        if alt:
            props.set("descr", alt)
        if name:
            props.set("name", name)
        return


def _set_arrow(connector: Any, tag: str, arrow: str) -> None:
    line_xml = connector._element.spPr.get_or_add_ln()
    qualified = "{http://schemas.openxmlformats.org/drawingml/2006/main}" + tag
    old = line_xml.find(qualified)
    if old is not None:
        line_xml.remove(old)
    if arrow != "none":
        endpoint = OxmlElement(f"a:{tag}")
        endpoint.set("type", arrow)
        line_xml.append(endpoint)


def _normalize_nested_zip(payload: bytes) -> bytes:
    """Normalize an embedded workbook so chart decks are reproducible too."""
    source = io.BytesIO(payload)
    destination = io.BytesIO()
    try:
        archive = zipfile.ZipFile(source, "r")
    except zipfile.BadZipFile:
        return payload
    with archive, zipfile.ZipFile(destination, "w", compression=zipfile.ZIP_DEFLATED) as output:
        seen: set[str] = set()
        for item in sorted(archive.infolist(), key=lambda value: value.filename):
            if item.filename in seen:
                raise UserInputError(f"嵌入式 Office 文件包含重复成员：{item.filename}")
            seen.add(item.filename)
            data = archive.read(item)
            if item.filename == "docProps/core.xml":
                data = CORE_DATE_RE.sub(rb"\g<1>2000-01-01T00:00:00Z\g<2>", data)
            normalized = zipfile.ZipInfo(item.filename, FIXED_TIMESTAMP)
            normalized.compress_type = zipfile.ZIP_DEFLATED
            normalized.external_attr = 0o600 << 16
            normalized.create_system = 3
            output.writestr(normalized, data)
    return destination.getvalue()


def _normalize_pptx(source: Path, destination: Path) -> None:
    with zipfile.ZipFile(source, "r") as archive, zipfile.ZipFile(
        destination, "w", compression=zipfile.ZIP_DEFLATED
    ) as output:
        seen: set[str] = set()
        for item in sorted(archive.infolist(), key=lambda value: value.filename):
            if item.filename in seen:
                raise UserInputError(f"PPTX 包含重复成员：{item.filename}")
            seen.add(item.filename)
            data = archive.read(item)
            if item.filename.startswith("ppt/embeddings/") and item.filename.endswith(".xlsx"):
                data = _normalize_nested_zip(data)
            elif item.filename.startswith("ppt/charts/chart") and item.filename.endswith(".xml"):
                data = _normalize_chart_axis_ids(data)
            normalized = zipfile.ZipInfo(item.filename, FIXED_TIMESTAMP)
            normalized.compress_type = zipfile.ZIP_DEFLATED
            normalized.external_attr = 0o600 << 16
            normalized.create_system = 3
            output.writestr(normalized, data)


class DeckBuilder:
    def __init__(self, spec: Mapping[str, Any], spec_path: Path):
        self.spec = spec
        self.spec_path = spec_path.resolve()
        self.base_dir = self.spec_path.parent
        self.theme = self._read_theme(spec.get("theme", {}))
        self.palette = self.theme["colors"]
        self.fonts = self.theme["fonts"]
        self.width, self.height = self._read_slide_size(spec.get("slide_size", "wide"))
        self.counts = {kind: 0 for kind in ("text", "shape", "line", "image", "table", "chart")}

    def _read_theme(self, value: Any) -> dict[str, dict[str, str]]:
        raw = require_mapping(value, "theme")
        _unknown_keys(raw, {"fonts", "colors"}, "theme")
        fonts = dict(DEFAULT_THEME["fonts"])
        if "fonts" in raw:
            supplied = require_mapping(raw["fonts"], "theme.fonts")
            _unknown_keys(supplied, {"heading", "body"}, "theme.fonts")
            for key, item in supplied.items():
                fonts[key] = _font_name(item, f"theme.fonts.{key}")
        colors = dict(DEFAULT_THEME["colors"])
        if "colors" in raw:
            supplied_colors = require_mapping(raw["colors"], "theme.colors")
            for key, item in supplied_colors.items():
                if not isinstance(key, str) or not key or len(key) > 100:
                    raise UserInputError("theme.colors 的颜色名必须是非空短字符串")
                colors[key] = color(item, f"theme.colors.{key}")
        return {"fonts": fonts, "colors": colors}

    def _read_slide_size(self, value: Any) -> tuple[float, float]:
        if isinstance(value, str):
            selected = choice(value, "slide_size", SLIDE_SIZES)
            return SLIDE_SIZES[selected]
        custom = require_mapping(value, "slide_size")
        _unknown_keys(custom, {"width", "height"}, "slide_size")
        width = number(custom.get("width"), "slide_size.width", minimum=1, maximum=56)
        height = number(custom.get("height"), "slide_size.height", minimum=1, maximum=56)
        return width, height

    def _rect(self, element: Mapping[str, Any], field: str) -> tuple[float, float, float, float]:
        x = number(element.get("x"), f"{field}.x", minimum=0)
        y = number(element.get("y"), f"{field}.y", minimum=0)
        width = number(element.get("w"), f"{field}.w", positive=True)
        height = number(element.get("h"), f"{field}.h", positive=True)
        if x + width > self.width + 0.001 or y + height > self.height + 0.001:
            raise UserInputError(
                f"{field} 越过幻灯片边界（画布 {self.width:.3f}×{self.height:.3f} 英寸）"
            )
        return x, y, width, height

    def _style_run(self, run: Any, style: Mapping[str, Any], field: str, *, default_size: float = 18) -> None:
        face = _font_name(style.get("font_face", self.fonts["body"]), f"{field}.font_face")
        size = number(style.get("font_size", default_size), f"{field}.font_size", positive=True, maximum=400)
        run.font.name = face
        _set_east_asian_font(run, face)
        run.font.size = Pt(size)
        if "bold" in style:
            run.font.bold = boolean(style["bold"], f"{field}.bold")
        if "italic" in style:
            run.font.italic = boolean(style["italic"], f"{field}.italic")
        rgb = color(style.get("color", "text"), f"{field}.color", self.palette)
        run.font.color.rgb = RGBColor.from_string(rgb)

    def _format_text_frame(
        self,
        frame: Any,
        element: Mapping[str, Any],
        field: str,
        *,
        default_size: float = 18,
    ) -> None:
        frame.clear()
        frame.word_wrap = True
        valign = choice(element.get("valign", "top"), f"{field}.valign", VERTICAL_ALIGNMENTS)
        frame.vertical_anchor = VERTICAL_ALIGNMENTS[valign]
        margin = element.get("margin", 0.08)
        if isinstance(margin, dict):
            _unknown_keys(margin, {"left", "right", "top", "bottom"}, f"{field}.margin")
            values = {
                side: number(margin.get(side, 0.08), f"{field}.margin.{side}", minimum=0, maximum=2)
                for side in ("left", "right", "top", "bottom")
            }
        else:
            uniform = number(margin, f"{field}.margin", minimum=0, maximum=2)
            values = {side: uniform for side in ("left", "right", "top", "bottom")}
        frame.margin_left = Inches(values["left"])
        frame.margin_right = Inches(values["right"])
        frame.margin_top = Inches(values["top"])
        frame.margin_bottom = Inches(values["bottom"])

        if "paragraphs" in element and "text" in element:
            raise UserInputError(f"{field} 不能同时包含 text 与 paragraphs")
        paragraphs = element.get("paragraphs", [element.get("text", "")])
        paragraphs = require_list(paragraphs, f"{field}.paragraphs", nonempty=True)
        if len(paragraphs) > MAX_PARAGRAPHS:
            raise UserInputError(f"{field}.paragraphs 不能超过 {MAX_PARAGRAPHS} 段")
        align_default = choice(element.get("align", "left"), f"{field}.align", ALIGNMENTS)
        for index, item in enumerate(paragraphs):
            paragraph_field = f"{field}.paragraphs[{index}]"
            if isinstance(item, str):
                paragraph_spec: Mapping[str, Any] = {"text": item}
            else:
                paragraph_spec = require_mapping(item, paragraph_field)
                _unknown_keys(paragraph_spec, PARAGRAPH_KEYS, paragraph_field)
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph_align = choice(
                paragraph_spec.get("align", align_default),
                f"{paragraph_field}.align",
                ALIGNMENTS,
            )
            paragraph.alignment = ALIGNMENTS[paragraph_align]
            if "level" in paragraph_spec:
                level = integer(paragraph_spec["level"], f"{paragraph_field}.level", minimum=0)
                if level > 8:
                    raise UserInputError(f"{paragraph_field}.level 不得大于 8")
                paragraph.level = level
            if "runs" in paragraph_spec and "text" in paragraph_spec:
                raise UserInputError(f"{paragraph_field} 不能同时包含 text 与 runs")
            run_specs = paragraph_spec.get("runs")
            if run_specs is None:
                run_specs = [{"text": paragraph_spec.get("text", "")}]
            run_specs = require_list(run_specs, f"{paragraph_field}.runs", nonempty=True)
            if len(run_specs) > MAX_RUNS_PER_PARAGRAPH:
                raise UserInputError(
                    f"{paragraph_field}.runs 不能超过 {MAX_RUNS_PER_PARAGRAPH} 个"
                )
            for run_index, run_item in enumerate(run_specs):
                run_field = f"{paragraph_field}.runs[{run_index}]"
                run_spec = require_mapping(run_item, run_field)
                _unknown_keys(run_spec, RUN_KEYS, run_field)
                text = optional_string(run_spec.get("text", ""), f"{run_field}.text")
                run = paragraph.add_run()
                run.text = text
                merged = dict(element)
                merged.update({key: value for key, value in paragraph_spec.items() if key not in {"runs", "text"}})
                merged.update(run_spec)
                self._style_run(run, merged, run_field, default_size=default_size)

    def _set_shape_fill(self, fill: Any, value: Any, field: str, default: str | None) -> None:
        selected = default if value is None else value
        if selected is None or selected == "none":
            fill.background()
            return
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(color(selected, field, self.palette))

    def _set_shape_line(self, line: Any, value: Any, field: str, default: str | None = None) -> None:
        if value is None:
            value = default
        if value is None or value == "none":
            line.fill.background()
            return
        if isinstance(value, dict):
            spec = value
            _unknown_keys(spec, LINE_STYLE_KEYS, field)
            selected = spec.get("color", "text")
            if "width" in spec:
                line.width = Pt(number(spec["width"], f"{field}.width", positive=True, maximum=50))
            if "dash" in spec:
                line.dash_style = DASHES[choice(spec["dash"], f"{field}.dash", DASHES)]
        else:
            spec = {}
            selected = value
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor.from_string(color(selected, f"{field}.color", self.palette))

    def _add_text(self, slide: Any, element: Mapping[str, Any], field: str) -> None:
        x, y, width, height = self._rect(element, field)
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        if "fill" in element:
            self._set_shape_fill(shape.fill, element["fill"], f"{field}.fill", None)
        if "line" in element:
            self._set_shape_line(shape.line, element["line"], f"{field}.line")
        self._format_text_frame(shape.text_frame, element, field)
        _set_alt_text(shape, element.get("alt"), element.get("name"))

    def _add_shape(self, slide: Any, element: Mapping[str, Any], field: str) -> None:
        x, y, width, height = self._rect(element, field)
        shape_name = element.get("shape_type", "rect")
        selected = choice(shape_name, f"{field}.shape_type", SHAPES)
        shape = slide.shapes.add_shape(
            SHAPES[selected], Inches(x), Inches(y), Inches(width), Inches(height)
        )
        self._set_shape_fill(shape.fill, element.get("fill"), f"{field}.fill", "primary")
        self._set_shape_line(shape.line, element.get("line"), f"{field}.line")
        if "rotation" in element:
            shape.rotation = number(element["rotation"], f"{field}.rotation", minimum=-360, maximum=360)
        if "text" in element or "paragraphs" in element:
            text_element = dict(element)
            text_element.setdefault("color", "white")
            text_element.setdefault("valign", "middle")
            text_element.setdefault("align", "center")
            self._format_text_frame(shape.text_frame, text_element, field)
        else:
            shape.text_frame.clear()
        _set_alt_text(shape, element.get("alt"), element.get("name"))

    def _add_line(self, slide: Any, element: Mapping[str, Any], field: str) -> None:
        x1 = number(element.get("x1"), f"{field}.x1", minimum=0, maximum=self.width)
        y1 = number(element.get("y1"), f"{field}.y1", minimum=0, maximum=self.height)
        x2 = number(element.get("x2"), f"{field}.x2", minimum=0, maximum=self.width)
        y2 = number(element.get("y2"), f"{field}.y2", minimum=0, maximum=self.height)
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
        )
        self._set_shape_line(connector.line, element.get("color", "text"), f"{field}.color")
        connector.line.width = Pt(number(element.get("width", 1.5), f"{field}.width", positive=True, maximum=50))
        connector.line.dash_style = DASHES[
            choice(element.get("dash", "solid"), f"{field}.dash", DASHES)
        ]
        arrows = ("none", "triangle", "stealth", "diamond", "oval", "arrow")
        begin = choice(element.get("begin_arrow", "none"), f"{field}.begin_arrow", arrows)
        end = choice(element.get("end_arrow", "none"), f"{field}.end_arrow", arrows)
        _set_arrow(connector, "headEnd", begin)
        _set_arrow(connector, "tailEnd", end)
        _set_alt_text(connector, element.get("alt"), element.get("name"))

    def _add_image(self, slide: Any, element: Mapping[str, Any], field: str) -> None:
        x, y, width, height = self._rect(element, field)
        image_path = safe_relative_file(
            self.base_dir, element.get("path"), field=f"{field}.path", must_exist=True
        )
        if image_path.stat().st_size > MAX_IMAGE_BYTES:
            raise UserInputError(f"{field}.path 图片文件超过 {MAX_IMAGE_BYTES} 字节")
        try:
            with warnings.catch_warnings():
                warnings.simplefilter("error", PILImage.DecompressionBombWarning)
                with PILImage.open(image_path) as image:
                    image_width, image_height = image.size
                    if image_width * image_height > MAX_IMAGE_PIXELS:
                        raise UserInputError(
                            f"{field}.path 图片像素数超过 {MAX_IMAGE_PIXELS}"
                        )
                    image.verify()
        except UserInputError:
            raise
        except Exception as exc:
            raise UserInputError(f"{field}.path 不是可用的栅格图片：{element.get('path')!r}") from exc
        if image_width <= 0 or image_height <= 0:
            raise UserInputError(f"{field}.path 的图片尺寸无效")
        fit = choice(element.get("fit", "contain"), f"{field}.fit", ("contain", "cover", "stretch"))
        image_ratio = image_width / image_height
        frame_ratio = width / height
        if fit == "contain":
            if image_ratio >= frame_ratio:
                target_width = width
                target_height = width / image_ratio
                target_x, target_y = x, y + (height - target_height) / 2
            else:
                target_height = height
                target_width = height * image_ratio
                target_x, target_y = x + (width - target_width) / 2, y
            picture = slide.shapes.add_picture(
                str(image_path), Inches(target_x), Inches(target_y), Inches(target_width), Inches(target_height)
            )
        else:
            picture = slide.shapes.add_picture(
                str(image_path), Inches(x), Inches(y), Inches(width), Inches(height)
            )
            if fit == "cover":
                if image_ratio > frame_ratio:
                    crop = (1 - frame_ratio / image_ratio) / 2
                    picture.crop_left = crop
                    picture.crop_right = crop
                elif image_ratio < frame_ratio:
                    crop = (1 - image_ratio / frame_ratio) / 2
                    picture.crop_top = crop
                    picture.crop_bottom = crop
        alt = element.get("alt")
        if alt is not None:
            alt = optional_string(alt, f"{field}.alt", maximum=1000)
        name = element.get("name")
        if name is not None:
            name = optional_string(name, f"{field}.name", maximum=200)
        _set_alt_text(picture, alt, name)

    def _write_cell(self, cell: Any, value: Any, field: str, *, header: bool, element: Mapping[str, Any]) -> None:
        cell.text = ""
        frame = cell.text_frame
        frame.clear()
        frame.margin_left = frame.margin_right = Inches(0.08)
        frame.margin_top = frame.margin_bottom = Inches(0.04)
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = frame.paragraphs[0]
        paragraph.alignment = ALIGNMENTS[
            choice(element.get("align", "left"), f"{field}.align", ALIGNMENTS)
        ]
        run = paragraph.add_run()
        run.text = _plain_cell(value, field)
        style = {
            "font_face": element.get("font_face", self.fonts["body"]),
            "font_size": element.get("font_size", 14),
            "color": element.get("header_color", "white") if header else element.get("color", "text"),
            "bold": header,
        }
        self._style_run(run, style, field, default_size=14)
        self._set_shape_fill(
            cell.fill,
            element.get("header_fill", "primary") if header else element.get("cell_fill", "white"),
            f"{field}.fill",
            None,
        )

    def _add_table(self, slide: Any, element: Mapping[str, Any], field: str) -> None:
        x, y, width, height = self._rect(element, field)
        rows = require_list(element.get("rows"), f"{field}.rows", nonempty=True)
        if len(rows) > 200:
            raise UserInputError(f"{field}.rows 不能超过 200 行")
        normalized: list[list[Any]] = []
        column_count: int | None = None
        for row_index, row in enumerate(rows):
            row_value = require_list(row, f"{field}.rows[{row_index}]", nonempty=True)
            if column_count is None:
                column_count = len(row_value)
                if column_count > 50:
                    raise UserInputError(f"{field}.rows 不能超过 50 列")
            if len(row_value) != column_count:
                raise UserInputError(f"{field}.rows[{row_index}] 的列数不一致")
            normalized.append(row_value)
        assert column_count is not None
        header_rows = integer(element.get("header_rows", 1), f"{field}.header_rows", minimum=0)
        if header_rows > len(rows):
            raise UserInputError(f"{field}.header_rows 不能超过表格行数")
        graphic = slide.shapes.add_table(
            len(rows), column_count, Inches(x), Inches(y), Inches(width), Inches(height)
        )
        table = graphic.table
        column_widths = element.get("column_widths")
        if column_widths is not None:
            values = require_list(column_widths, f"{field}.column_widths", nonempty=True)
            if len(values) != column_count:
                raise UserInputError(f"{field}.column_widths 数量必须等于列数")
            parsed = [
                number(value, f"{field}.column_widths[{index}]", positive=True)
                for index, value in enumerate(values)
            ]
            if abs(sum(parsed) - width) > 0.02:
                raise UserInputError(f"{field}.column_widths 的总和必须等于 w")
            for index, value in enumerate(parsed):
                table.columns[index].width = Inches(value)
        for row_index, row in enumerate(normalized):
            table.rows[row_index].height = Inches(height / len(normalized))
            for column_index, value in enumerate(row):
                self._write_cell(
                    table.cell(row_index, column_index),
                    value,
                    f"{field}.rows[{row_index}][{column_index}]",
                    header=row_index < header_rows,
                    element=element,
                )
        _set_alt_text(graphic, element.get("alt"), element.get("name"))

    def _chart_colors(self, element: Mapping[str, Any], field: str) -> list[str]:
        defaults = ["primary", "secondary", "accent", "muted"]
        values = element.get("colors", defaults)
        values = require_list(values, f"{field}.colors", nonempty=True)
        return [color(value, f"{field}.colors[{index}]", self.palette) for index, value in enumerate(values)]

    def _add_chart(self, slide: Any, element: Mapping[str, Any], field: str) -> None:
        x, y, width, height = self._rect(element, field)
        chart_name = choice(element.get("chart"), f"{field}.chart", CHARTS)
        categories = require_list(element.get("categories"), f"{field}.categories", nonempty=True)
        if len(categories) > 1000:
            raise UserInputError(f"{field}.categories 不能超过 1000 项")
        chart_data = ChartData()
        chart_data.categories = [_plain_cell(value, f"{field}.categories") for value in categories]
        series_values = require_list(element.get("series"), f"{field}.series", nonempty=True)
        if len(series_values) > MAX_CHART_SERIES:
            raise UserInputError(f"{field}.series 不能超过 {MAX_CHART_SERIES} 个")
        if chart_name == "pie" and len(series_values) != 1:
            raise UserInputError(f"{field}.series：饼图必须且只能有一个数据系列")
        for series_index, series_item in enumerate(series_values):
            series_field = f"{field}.series[{series_index}]"
            series = require_mapping(series_item, series_field)
            _unknown_keys(series, CHART_SERIES_KEYS, series_field)
            name = optional_string(series.get("name", f"系列 {series_index + 1}"), f"{series_field}.name", maximum=200)
            values = require_list(series.get("values"), f"{series_field}.values")
            if len(values) != len(categories):
                raise UserInputError(f"{series_field}.values 数量必须等于 categories 数量")
            parsed_values: list[float | None] = []
            for value_index, value in enumerate(values):
                if value is None and chart_name != "pie":
                    parsed_values.append(None)
                else:
                    parsed_values.append(number(value, f"{series_field}.values[{value_index}]"))
            chart_data.add_series(name, parsed_values)
        frame = slide.shapes.add_chart(
            CHARTS[chart_name], Inches(x), Inches(y), Inches(width), Inches(height), chart_data
        )
        chart = frame.chart
        if "style" in element:
            chart_style = integer(element["style"], f"{field}.style", minimum=1)
            if chart_style > 48:
                raise UserInputError(f"{field}.style 不得大于 48")
            chart.chart_style = chart_style
        title = element.get("title")
        if title is not None:
            chart.has_title = True
            chart.chart_title.text_frame.text = optional_string(title, f"{field}.title", maximum=500)
        legend = element.get("legend")
        if legend is None:
            legend = "bottom" if len(series_values) > 1 else "none"
        legend = choice(legend, f"{field}.legend", ("none", *LEGENDS.keys()))
        chart.has_legend = legend != "none"
        if legend != "none":
            chart.legend.position = LEGENDS[legend]
            chart.legend.include_in_layout = False
        if "data_labels" in element:
            plot = chart.plots[0]
            plot.has_data_labels = boolean(element["data_labels"], f"{field}.data_labels")
        colors = self._chart_colors(element, field)
        if chart_name == "pie":
            for index, point in enumerate(chart.series[0].points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(colors[index % len(colors)])
        elif chart_name == "line":
            for index, series in enumerate(chart.series):
                series.format.line.color.rgb = RGBColor.from_string(colors[index % len(colors)])
                series.format.line.width = Pt(2.25)
        else:
            for index, series in enumerate(chart.series):
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = RGBColor.from_string(colors[index % len(colors)])
        _set_alt_text(frame, element.get("alt"), element.get("name"))

    def _add_element(self, slide: Any, value: Any, slide_index: int, element_index: int) -> None:
        field = f"slides[{slide_index}].elements[{element_index}]"
        element = require_mapping(value, field)
        if "name" in element:
            optional_string(element["name"], f"{field}.name", maximum=200)
        if "alt" in element:
            optional_string(element["alt"], f"{field}.alt", maximum=1000)
        kind = choice(element.get("type"), f"{field}.type", self.counts)
        _unknown_keys(element, ELEMENT_KEYS[kind], field)
        dispatch = {
            "text": self._add_text,
            "shape": self._add_shape,
            "line": self._add_line,
            "image": self._add_image,
            "table": self._add_table,
            "chart": self._add_chart,
        }
        dispatch[kind](slide, element, field)
        self.counts[kind] += 1

    def build(self) -> Presentation:
        _unknown_keys(
            self.spec,
            {"schema_version", "route", "profile", "meta", "slide_size", "theme", "slides"},
            "spec",
        )
        version = self.spec.get("schema_version")
        if version != SCHEMA_VERSION:
            raise UserInputError(f"schema_version 必须为 {SCHEMA_VERSION}")
        route = choice(self.spec.get("route", "generate"), "route", ROUTES)
        if route not in BUILDER_ROUTES:
            raise UserInputError(
                "bundled builder 仅支持 route=generate，或为 create-template 生成可选评审 deck；"
                "fill-template/enhance 必须使用能保留原生 PPTX 的运行时"
            )
        if "profile" in self.spec:
            if route != "generate":
                raise UserInputError("profile 仅允许用于 route=generate")
            choice(self.spec["profile"], "profile", PROFILES)
        slides = require_list(self.spec.get("slides"), "slides", nonempty=True)
        if len(slides) > 500:
            raise UserInputError("slides 不能超过 500 页")

        presentation = Presentation()
        presentation.slide_width = Inches(self.width)
        presentation.slide_height = Inches(self.height)
        properties = presentation.core_properties
        properties.created = datetime(2000, 1, 1)
        properties.modified = datetime(2000, 1, 1)
        properties.revision = 1
        meta = require_mapping(self.spec.get("meta", {}), "meta")
        allowed_meta = {
            "title",
            "subject",
            "author",
            "keywords",
            "comments",
            "category",
            "content_status",
            "identifier",
            "language",
            "version",
        }
        _unknown_keys(meta, allowed_meta, "meta")
        for key, value in meta.items():
            setattr(properties, key, optional_string(value, f"meta.{key}", maximum=255))
        properties.last_modified_by = properties.author or "ppt-master"

        blank_layout = presentation.slide_layouts[6]
        for slide_index, slide_item in enumerate(slides):
            field = f"slides[{slide_index}]"
            slide_spec = require_mapping(slide_item, field)
            _unknown_keys(slide_spec, {"name", "background", "notes", "elements"}, field)
            slide = presentation.slides.add_slide(blank_layout)
            if "name" in slide_spec:
                slide._element.cSld.set(
                    "name", optional_string(slide_spec["name"], f"{field}.name", maximum=200)
                )
            background = color(
                slide_spec.get("background", "background"), f"{field}.background", self.palette
            )
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string(background)
            if "notes" in slide_spec:
                slide.notes_slide.notes_text_frame.text = optional_string(
                    slide_spec["notes"], f"{field}.notes", maximum=50000
                )
            elements = require_list(slide_spec.get("elements", []), f"{field}.elements")
            if len(elements) > 1000:
                raise UserInputError(f"{field}.elements 不能超过 1000 个")
            for element_index, element in enumerate(elements):
                self._add_element(slide, element, slide_index, element_index)
        return presentation


def _safe_output_path(output: Path) -> Path:
    expanded = output.expanduser()
    if expanded.is_symlink():
        raise UserInputError(f"输出路径不能是符号链接：{expanded}")
    return expanded.parent.resolve(strict=False) / expanded.name


def save_deterministic(
    presentation: Presentation, output: Path, *, force: bool = False
) -> Path:
    output = _safe_output_path(output)
    if output.exists() and output.is_dir():
        raise UserInputError(f"输出路径是目录：{output}")
    if output.exists() and not force:
        raise UserInputError(
            f"输出文件已存在，为避免覆盖已停止：{output}；确认替换生成物时使用 --force"
        )
    output.parent.mkdir(parents=True, exist_ok=True)
    first_fd, first_name = tempfile.mkstemp(prefix=f".{output.name}.", suffix=".raw", dir=output.parent)
    second_fd, second_name = tempfile.mkstemp(prefix=f".{output.name}.", suffix=".normalized", dir=output.parent)
    os.close(first_fd)
    os.close(second_fd)
    raw_path, normalized_path = Path(first_name), Path(second_name)
    try:
        presentation.save(raw_path)
        _normalize_pptx(raw_path, normalized_path)
        os.replace(normalized_path, output)
    finally:
        for path in (raw_path, normalized_path):
            try:
                path.unlink()
            except FileNotFoundError:
                pass
    return output


def build_from_file(spec_path: Path, output: Path, *, force: bool = False) -> dict[str, Any]:
    spec_path = spec_path.expanduser().resolve()
    if not spec_path.is_file():
        raise UserInputError(f"spec 文件不存在：{spec_path}")
    safe_output = _safe_output_path(output)
    same_input = safe_output == spec_path
    if safe_output.exists() and not same_input:
        try:
            same_input = os.path.samefile(safe_output, spec_path)
        except OSError:
            same_input = False
    if same_input:
        raise UserInputError("输出路径不能与 Deck Spec 输入文件相同")
    if safe_output.exists() and safe_output.is_dir():
        raise UserInputError(f"输出路径是目录：{safe_output}")
    if safe_output.exists() and not force:
        raise UserInputError(
            f"输出文件已存在，为避免覆盖已停止：{safe_output}；确认替换生成物时使用 --force"
        )
    spec = load_json_object(spec_path)
    builder = DeckBuilder(spec, spec_path)
    presentation = builder.build()
    written_output = save_deterministic(presentation, safe_output, force=force)
    return {
        "ok": True,
        "spec": str(spec_path),
        "output": str(written_output),
        "slides": len(presentation.slides),
        "elements": dict(builder.counts),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="从 JSON spec 构建原生可编辑 PPTX")
    parser.add_argument("spec", type=Path, help="JSON spec 文件")
    parser.add_argument("-o", "--output", type=Path, required=True, help="输出 .pptx 文件")
    parser.add_argument(
        "--force", action="store_true", help="确认覆盖已存在的常规输出文件；仍拒绝符号链接"
    )
    args = parser.parse_args(argv)
    if args.output.suffix.lower() != ".pptx":
        parser.exit(2, "错误：输出文件必须使用 .pptx 扩展名\n")
    try:
        result = build_from_file(args.spec, args.output, force=args.force)
    except UserInputError as exc:
        parser.exit(2, f"错误：{exc}\n")
    except (OSError, zipfile.BadZipFile) as exc:
        parser.exit(1, f"构建失败：{exc}\n")
    print(json.dumps(result, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
